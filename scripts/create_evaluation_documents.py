"""Create public, deterministic format fixtures; requires the documents extra."""

from pathlib import Path


def create(directory):
    from docx import Document
    from docx.oxml import OxmlElement
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen import canvas
    from PIL import Image, ImageDraw, ImageFont

    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("DejaVuSans.ttf", 32)
    draw.text((80, 100), "Orion emergency recovery", font=font, fill="black")
    draw.text((80, 200), "The recovery code is DR-4826.", font=font, fill="black")
    draw.text(
        (80, 300), "The recovery time objective is 45 minutes.", font=font, fill="black"
    )
    image_path = directory / "recovery.png"
    image.save(image_path)

    pdf = canvas.Canvas(str(directory / "text.pdf"))
    pdf.drawString(50, 780, "Orion backup retention is 35 days.")
    pdf.showPage()
    pdf.drawString(50, 780, "Backup restoration drills happen every 14 days.")
    pdf.save()
    image.save(directory / "scan.pdf", "PDF", resolution=150)

    pdf = canvas.Canvas(str(directory / "mixed.pdf"))
    pdf.drawString(50, 780, "Orion capacity planning: the reserve is 27 percent.")
    pdf.drawString(50, 720, "Service        Daily capacity")
    pdf.drawString(50, 690, "Atlas          840 requests")
    pdf.drawString(50, 660, "Boreal         260 requests")
    for y in (735, 705, 675, 645):
        pdf.line(45, y, 360, y)
    for x in (45, 165, 360):
        pdf.line(x, 645, x, 735)
    pdf.drawString(50, 580, "Capacity formula: C = workers * 42.")
    pdf.drawImage(str(image_path), 50, 110, width=480, height=280)
    pdf.save()

    doc = Document()
    doc.add_heading("Orion service handbook", 0)
    doc.add_paragraph("A production incident must be acknowledged within 7 minutes.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Severity", "Response"
    table.cell(1, 0).text, table.cell(1, 1).text = "P1", "7 minutes"
    doc.add_picture(str(image_path))
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "C = workers * 42"
    run.append(text)
    math.append(run)
    doc.add_paragraph()._p.append(math)
    doc.save(directory / "handbook.docx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Orion release policy"
    slide.placeholders[1].text = "Production releases happen on Tuesday at 09:30 UTC."
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Release verification"
    slide.placeholders[1].text = "The rollback window is 18 minutes."
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(3), width=Inches(5))
    presentation.save(directory / "release.pptx")
    return [
        directory / name
        for name in (
            "text.pdf",
            "scan.pdf",
            "mixed.pdf",
            "handbook.docx",
            "release.pptx",
        )
    ]


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    for path in create(args.directory):
        print(path.name)
